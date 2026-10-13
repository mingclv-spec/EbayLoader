"""Book-to-eBay listing draft UI. Run: streamlit run ebayloader/webapp.py."""
from __future__ import annotations
import csv, io, json, os
import requests
from datetime import datetime
from typing import Any, Dict
import streamlit as st
from PIL import Image
from ebayloader.book_research import extract_isbns, research_isbn, research_links, web_search

st.set_page_config(page_title="EbayLoader · Book Listing Assistant", page_icon="📚", layout="wide")
st.title("📚 Book → eBay Listing Assistant")
st.caption("Photograph a book, research the edition, review listing fields, and export CSV or Excel.")
st.info("This tool prepares drafts only; it does not publish to eBay. Verify every bibliographic detail and condition against the physical copy.")

DEFAULT_FIELDS = {"sku":"","title":"","subtitle":"","author":"","publisher":"","publication_date":"","edition":"","isbn":"","format":"Paperback","language":"English","pages":"","condition":"Good","condition_notes":"","description":"","category_id":"261186","price":"","quantity":"1","location":"","shipping_cost":"","return_days":"30","photo_filenames":"","research_notes":""}

def txt(value: Any) -> str:
    if value is None: return ""
    if isinstance(value,(list,dict)): return json.dumps(value,ensure_ascii=False)
    return str(value)

def csv_bytes(row: Dict[str,Any]) -> bytes:
    out=io.StringIO(newline="")
    writer=csv.DictWriter(out,fieldnames=list(row.keys()))
    writer.writeheader(); writer.writerow({k:txt(v) for k,v in row.items()})
    return out.getvalue().encode("utf-8-sig")

def xlsx_bytes(row: Dict[str,Any]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb=Workbook(); ws=wb.active; ws.title="eBay Listing"
    ws.append(list(row.keys())); ws.append([txt(v) for v in row.values()])
    for cell in ws[1]:
        cell.font=Font(bold=True,color="FFFFFF"); cell.fill=PatternFill("solid",fgColor="25476A")
    ws.freeze_panes="A2"; ws.auto_filter.ref=ws.dimensions
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width=min(max(max(len(str(c.value or "")) for c in col)+2,14),42)
    out=io.BytesIO(); wb.save(out); return out.getvalue()

def ocr_photos(images):
    try: import pytesseract
    except ImportError: return "", "Optional pytesseract package is missing; manual ISBN entry is available."
    chunks=[]
    for item in images:
        try: chunks.append(pytesseract.image_to_string(Image.open(item).convert("RGB")))
        except Exception as exc: chunks.append("OCR error for "+item.name+": "+str(exc))
    return "\n".join(chunks), ""

with st.sidebar:
    st.header("Workflow")
    st.markdown("1. Upload cover, ISBN/barcode, copyright page and defects\n2. Identify ISBN\n3. Research metadata and sources\n4. Edit listing fields\n5. Export CSV / Excel")
    st.caption("OCR is optional and requires the Python package plus the Tesseract engine. Enter ISBN manually if OCR is unavailable.")
    st.caption("eBay sold comparisons open search results for manual review; this app does not claim to fetch verified sold prices.")

images=st.file_uploader("Book photos",type=["jpg","jpeg","png","webp"],accept_multiple_files=True)
photo_names=";".join(x.name for x in images) if images else ""
if images:
    cols=st.columns(min(4,len(images)))
    for i,item in enumerate(images):
        with cols[i%len(cols)]: st.image(item,caption=item.name,use_container_width=True)
if "draft" not in st.session_state: st.session_state.draft=DEFAULT_FIELDS.copy()
if "research" not in st.session_state: st.session_state.research={}
if "research_links" not in st.session_state: st.session_state.research_links=[]
if "web_results" not in st.session_state: st.session_state.web_results=[]

if images and st.button("Read ISBN from photos (optional OCR)"):
    with st.spinner("Reading photo text…"):
        text,note=ocr_photos(images)
        st.session_state.ocr_text=text; st.session_state.ocr_note=note
if st.session_state.get("ocr_text"):
    with st.expander("OCR text (review; recognition may be wrong)"):
        st.text_area("Recognised text",st.session_state.ocr_text,height=120)
        st.caption(st.session_state.get("ocr_note",""))
matches=extract_isbns(st.session_state.get("ocr_text",""))
if matches: st.success("Possible ISBN(s): "+", ".join(matches))

st.subheader("1 · Identify the book")
draft=st.session_state.draft
with st.form("research_form"):
    isbn=st.text_input("ISBN-10 or ISBN-13",value=draft.get("isbn") or (matches[0] if matches else ""))
    manual_title=st.text_input("Title (if ISBN unavailable)",value=draft.get("title",""))
    manual_author=st.text_input("Author(s)",value=draft.get("author",""))
    go=st.form_submit_button("Research book")
if go:
    data={}
    if isbn.strip():
        try: data=research_isbn(isbn)
        except Exception as exc: st.error(str(exc))
    else: st.warning("No ISBN supplied. Manual title search is less reliable for edition matching.")
    meta=data.get("metadata",{}); merged=st.session_state.draft.copy()
    merged.update({"isbn":data.get("isbn",isbn.strip()),"title":meta.get("title") or manual_title,"subtitle":meta.get("subtitle") or merged.get("subtitle",""),"author":txt(meta.get("authors") or meta.get("author_keys") or manual_author),"publisher":txt(meta.get("publisher") or meta.get("publishers") or merged.get("publisher","")),"publication_date":txt(meta.get("publication_date") or merged.get("publication_date","")),"pages":txt(meta.get("pages") or merged.get("pages","")),"language":txt(meta.get("language") or merged.get("language","")),"description":txt(meta.get("summary") or merged.get("description","")),"photo_filenames":photo_names})
    st.session_state.draft=merged; st.session_state.research=data
    query=" ".join(x for x in (merged.get("title"),merged.get("author"),merged.get("isbn")) if x)
    st.session_state.web_results=web_search(query) if query else []
    st.session_state.research_links=research_links(merged.get("title",""),merged.get("author",""),merged.get("isbn",""))
    st.rerun()

research=st.session_state.research
if research:
    st.subheader("2 · Research evidence")
    for src in research.get("sources",[]):
        st.markdown("- **"+src["source"]+"** — "+src["status"]+": "+src["note"]+" ([open source]("+src["url"]+"))")
    if research.get("metadata"):
        with st.expander("Raw metadata returned by sources"): st.json(research["metadata"])
if st.session_state.research_links:
    st.markdown("**Publisher, web, and eBay comparison links**")
    for item in st.session_state.research_links: st.markdown("- ["+item["source"]+"]("+item["url"]+") — "+item["note"])
if st.session_state.web_results:
    with st.expander("Best-effort web search results"):
        for item in st.session_state.web_results:
            if item.get("url"): st.markdown("- ["+(item.get("title") or item["url"])+"]("+item["url"]+")")

st.subheader("3 · Review and edit listing fields")
st.caption("Confirm edition, item specifics, condition and price. Unknown facts should remain blank.")
d=st.session_state.draft
formats=["Paperback","Hardcover","Board book","Spiral-bound","Loose-leaf","Audiobook","Other"]
conditions=["New","Like New","Very Good","Good","Acceptable"]
with st.form("listing_form"):
    a,b,c=st.columns(3)
    with a:
        sku=st.text_input("SKU",d.get("sku","")); title=st.text_input("eBay title",d.get("title",""))
        subtitle=st.text_input("Subtitle",d.get("subtitle","")); author=st.text_input("Author",d.get("author",""))
        publisher=st.text_input("Publisher",d.get("publisher","")); edition=st.text_input("Edition",d.get("edition",""))
        isbn2=st.text_input("ISBN",d.get("isbn",""))
    with b:
        pubdate=st.text_input("Publication date",d.get("publication_date",""))
        fmt=st.selectbox("Format",formats,index=formats.index(d.get("format")) if d.get("format") in formats else 0)
        language=st.text_input("Language",d.get("language","English")); pages=st.text_input("Page count",d.get("pages",""))
        condition=st.selectbox("Condition",conditions,index=conditions.index(d.get("condition")) if d.get("condition") in conditions else 3)
        condition_notes=st.text_input("Condition notes",d.get("condition_notes","")); category=st.text_input("eBay category ID",d.get("category_id","261186"))
    with c:
        price=st.text_input("Price (AUD)",d.get("price","")); quantity=st.text_input("Quantity",d.get("quantity","1"))
        location=st.text_input("Item location",d.get("location","")); shipping=st.text_input("Shipping cost (AUD)",d.get("shipping_cost",""))
        returns=st.text_input("Return days",d.get("return_days","30")); photos=st.text_input("Photo filenames",photo_names or d.get("photo_filenames",""))
    description=st.text_area("Listing description",d.get("description",""),height=160)
    notes=st.text_area("Research notes / missing information",d.get("research_notes",""))
    save=st.form_submit_button("Save listing draft")
if save:
    st.session_state.draft={"sku":sku,"title":title,"subtitle":subtitle,"author":author,"publisher":publisher,"publication_date":pubdate,"edition":edition,"isbn":isbn2,"format":fmt,"language":language,"pages":pages,"condition":condition,"condition_notes":condition_notes,"description":description,"category_id":category,"price":price,"quantity":quantity,"location":location,"shipping_cost":shipping,"return_days":returns,"photo_filenames":photos,"research_notes":notes}
    st.success("Draft saved in this browser session.")

st.subheader("AI-assisted description")
st.caption("Optional: uses your own OpenAI API key. It drafts copy only from the fields you have confirmed; review before using.")
if st.button("Generate listing description with AI"):
    api_key = os.environ.get("OPENAI_API_KEY", "").strip()
    if not api_key:
        st.warning("Set OPENAI_API_KEY in your environment and restart the app to enable AI writing.")
    else:
        facts = {k: st.session_state.draft.get(k, "") for k in ("title", "subtitle", "author", "publisher", "publication_date", "edition", "isbn", "format", "language", "pages", "condition", "condition_notes")}
        prompt = (
            "Write a concise, buyer-friendly second-hand book listing description. "
            "Use only facts in the supplied JSON. Do not invent contents, edition details, "
            "included supplements, access codes, signatures, provenance, or guarantees. "
            "If a value is blank, omit it. Mention the supplied condition notes neutrally. "
            "Return plain text only, without a price or claims that the book is rare.\\n\\n"
            + json.dumps(facts, ensure_ascii=False)
        )
        try:
            response = requests.post(
                "https://api.openai.com/v1/chat/completions",
                headers={"Authorization": "Bearer " + api_key, "Content-Type": "application/json"},
                json={"model": os.environ.get("OPENAI_MODEL", "gpt-4.1-mini"), "messages": [{"role": "user", "content": prompt}], "temperature": 0.2},
                timeout=45,
            )
            response.raise_for_status()
            description = response.json()["choices"][0]["message"]["content"].strip()
            st.session_state.draft["description"] = description
            st.success("AI description generated. It has been placed in the draft; review and edit it in the description field above.")
            st.rerun()
        except Exception as exc:
            st.error("AI generation failed: " + str(exc)[:300])

st.subheader("4 · Export")
row=st.session_state.draft.copy()
row["research_sources"]=json.dumps(research.get("sources",[]),ensure_ascii=False)
row["research_links"]=json.dumps(st.session_state.research_links,ensure_ascii=False)
row["generated_at_utc"]=datetime.utcnow().replace(microsecond=0).isoformat()+"Z"
if not row.get("title"): st.warning("Add a title before exporting.")
x,y=st.columns(2)
with x: st.download_button("Download CSV",csv_bytes(row),"ebay_book_listing.csv","text/csv",disabled=not bool(row.get("title")))
with y: st.download_button("Download Excel (.xlsx)",xlsx_bytes(row),"ebay_book_listing.xlsx","application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",disabled=not bool(row.get("title")))
with st.expander("Customise listing fields"):
    st.markdown("Add, remove or rename keys in DEFAULT_FIELDS near the top of ebayloader/webapp.py. The exported record includes source URLs to help verify claims.")
